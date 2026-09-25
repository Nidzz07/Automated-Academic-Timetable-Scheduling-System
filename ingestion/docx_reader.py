"""Read every paragraph and table of a .docx in true document order.

This module is the shared foundation for the class, lab and faculty timetable
parsers. It knows about WordprocessingML structure and nothing about
timetables: no slots, no subjects, no field orders. Those belong to the
per-source parsers built on top of it.

Why not python-docx's convenience API
-------------------------------------
Most timetable grids in ``data/real/`` sit inside Word content controls
(``<w:sdt>``). ``Document.tables`` and ``Document.paragraphs`` only look at
the direct children of ``<w:body>`` and silently skip everything inside an
``<w:sdt>``. On the EVEN class timetable that is 4 tables seen out of 21. A
parser built on them would read about a fifth of the data and report success.
So this module walks the XML itself, and ``Document.tables`` /
``Document.paragraphs`` are not used anywhere.

Logical grid
------------
A Word table has a physical grid (``<w:tblGrid>``) and cells that span several
grid columns (``<w:gridSpan>``) or continue a cell from the row above
(``<w:vMerge>``). The EVEN class tables declare 42-44 grid columns, yet every
row puts its cells at the same few grid offsets: the table is really 6 columns
(Time + Monday..Friday) drawn on a finer grid. Expanding ``gridSpan`` alone
would give 42 columns of repeated text.

The logical columns are therefore the intervals between *cell boundaries* -
the union, over all rows, of the grid offsets where some cell starts. Every
boundary that any row actually uses is kept, so no real column is lost, and
grid lines no cell starts on disappear. Columns are never merged by comparing
their text: two genuinely different columns can hold identical text.

One geometric exception. Where a row's cell edge was dragged a hair off the
edge in the rows around it, the union produces a *sliver* column: the BE table
in the EVEN class file has one 105 twips (~2 mm) wide, next to columns of
2000+. No text fits in 2 mm, so a logical column narrower than
:data:`SLIVER_TWIPS` is treated as edge misalignment: the less-used of its two
boundaries is dropped and every cell edge snaps to the nearest remaining
boundary by physical distance. Across all four source files the narrowest
column that survives is 720 twips, so the threshold sits in a wide empty gap.
Genuine sub-columns (the ODD SE tables split Tuesday and Wednesday into two
parallel columns of 1000-2000 twips each) are untouched.
"""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import docx
from docx.document import Document
from docx.oxml.ns import qn
from lxml.etree import _Element

BlockKind = Literal["paragraph", "table"]

_W_P = qn("w:p")
_W_TBL = qn("w:tbl")
_W_TR = qn("w:tr")
_W_TC = qn("w:tc")
_W_T = qn("w:t")
_W_BR = qn("w:br")
_W_CR = qn("w:cr")
_W_TAB = qn("w:tab")
_W_VAL = qn("w:val")

# Containers that wrap block- or cell-level content without being content
# themselves. Their children are walked as if they sat in the parent.
_W_SDT = qn("w:sdt")
_W_SDT_CONTENT = qn("w:sdtContent")
_W_CUSTOM_XML = qn("w:customXml")
_TRANSPARENT = {_W_SDT_CONTENT, _W_CUSTOM_XML}

# A logical column narrower than this (a quarter inch) is edge misalignment,
# not a column. See the module docstring.
SLIVER_TWIPS = 360


def open_document(path: str | Path) -> Document:
    """Open a .docx. A thin wrapper so callers need not import python-docx."""
    return docx.Document(str(path))


# --------------------------------------------------------------------------- #
# Document-order walk
# --------------------------------------------------------------------------- #


def _children(element: _Element, wanted: set[str]) -> Iterator[_Element]:
    """Yield children of ``element`` whose tag is in ``wanted``, in order.

    Descends through ``<w:sdt>`` (into its ``<w:sdtContent>`` only - its
    ``<w:sdtPr>`` properties hold no document content) and ``<w:customXml>``.
    """
    for child in element:
        if child.tag in wanted:
            yield child
        elif child.tag == _W_SDT:
            content = child.find(_W_SDT_CONTENT)
            if content is not None:
                yield from _children(content, wanted)
        elif child.tag in _TRANSPARENT:
            yield from _children(child, wanted)


def walk_body(document: Document) -> Iterator[tuple[BlockKind, _Element]]:
    """Yield ``("paragraph" | "table", element)`` for every body block, in order.

    Blocks wrapped in content controls are yielded at their true position.
    Tables are yielded whole - paragraphs inside table cells are not yielded
    here; read them through :func:`logical_grid`.
    """
    for element in _children(document.element.body, {_W_P, _W_TBL}):
        yield ("paragraph" if element.tag == _W_P else "table"), element


def paragraph_text(paragraph: _Element) -> str:
    """Text of one ``<w:p>``: runs concatenated, ``<w:br>``/``<w:cr>`` as newlines.

    Tabs become ``\\t``. Nothing is stripped - whitespace decisions belong to
    the parsers.
    """
    parts: list[str] = []
    for node in paragraph.iter(_W_T, _W_BR, _W_CR, _W_TAB):
        if node.tag == _W_T:
            parts.append(node.text or "")
        elif node.tag == _W_TAB:
            parts.append("\t")
        else:
            parts.append("\n")
    return "".join(parts)


def cell_text(cell: _Element) -> str:
    """Text of one ``<w:tc>``, one line per paragraph, joined with ``\\n``.

    Lines are kept intact: a lab cell listing four parallel batch sessions
    comes back as four lines, not one run-together string.
    """
    return "\n".join(paragraph_text(p) for p in _children(cell, {_W_P}))


# --------------------------------------------------------------------------- #
# Logical grid
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class GridCell:
    """One physical ``<w:tc>`` placed on the logical grid.

    ``row``/``col`` are the top-left logical position; ``row_span``/``col_span``
    how many logical rows/columns it covers once ``vMerge`` continuations and
    ``gridSpan`` are resolved.
    """

    row: int
    col: int
    row_span: int
    col_span: int
    text: str


@dataclass
class _RawCell:
    row: int
    grid_start: int
    grid_end: int
    vmerge: Literal["restart", "continue"] | None
    text: str


@dataclass
class _Placed:
    row: int
    row_end: int
    col: int
    col_end: int
    text: str


def _prop_val(cell: _Element, name: str) -> str | None:
    """Value of ``w:tcPr/w:<name>/@w:val``; ``""`` if the element has no val."""
    props = cell.find(qn("w:tcPr"))
    if props is None:
        return None
    node = props.find(qn(name))
    if node is None:
        return None
    return node.get(_W_VAL, "")


def _row_grid_skip(row: _Element, name: str) -> int:
    """``w:trPr/w:gridBefore`` or ``w:gridAfter`` - grid columns a row leaves empty."""
    props = row.find(qn("w:trPr"))
    node = props.find(qn(name)) if props is not None else None
    return int(node.get(_W_VAL, "0")) if node is not None else 0


def _raw_cells(table: _Element) -> list[_RawCell]:
    """Every cell with its physical grid interval ``[grid_start, grid_end)``."""
    cells: list[_RawCell] = []
    for r, row in enumerate(_children(table, {_W_TR})):
        offset = _row_grid_skip(row, "w:gridBefore")
        for cell in _children(row, {_W_TC}):
            span = int(_prop_val(cell, "w:gridSpan") or "1")
            vmerge_val = _prop_val(cell, "w:vMerge")
            # A bare <w:vMerge/> means "continue"; val="restart" starts a merge.
            vmerge: Literal["restart", "continue"] | None
            if vmerge_val is None:
                vmerge = None
            elif vmerge_val == "restart":
                vmerge = "restart"
            else:
                vmerge = "continue"
            cells.append(_RawCell(r, offset, offset + span, vmerge, cell_text(cell)))
            offset += span
    return cells


def _grid_positions(table: _Element, right_edge: int) -> list[float]:
    """Physical x-position (twips) of every grid offset ``0..right_edge``.

    Falls back to one unit per grid column when ``<w:tblGrid>`` is missing or
    shorter than the rows need; sliver folding is then effectively disabled.
    """
    widths = [
        float(col.get(qn("w:w")) or 0)
        for col in table.findall(f"{qn('w:tblGrid')}/{qn('w:gridCol')}")
    ]
    if len(widths) < right_edge:
        widths = [float(SLIVER_TWIPS)] * right_edge
    positions = [0.0]
    for width in widths[:right_edge]:
        positions.append(positions[-1] + width)
    return positions


def _logical_boundaries(raw: list[_RawCell], right_edge: int, position: list[float]) -> list[int]:
    """Grid offsets that delimit logical columns, with sliver columns folded."""
    starts: dict[int, int] = {}
    for c in raw:
        starts[c.grid_start] = starts.get(c.grid_start, 0) + 1
    boundaries = sorted(set(starts) | {right_edge})
    while len(boundaries) > 2:
        widths = [
            position[boundaries[i + 1]] - position[boundaries[i]]
            for i in range(len(boundaries) - 1)
        ]
        narrowest = min(range(len(widths)), key=widths.__getitem__)
        if widths[narrowest] >= SLIVER_TWIPS:
            break
        # Drop whichever edge of the sliver fewer cells start on; the outer
        # edges of the table are never dropped.
        droppable = [i for i in (narrowest, narrowest + 1) if 0 < i < len(boundaries) - 1]
        drop = min(droppable, key=lambda i: (starts.get(boundaries[i], 0), -i))
        del boundaries[drop]
    return boundaries


def logical_cells(table: _Element) -> tuple[list[GridCell], int, int]:
    """Resolve a ``<w:tbl>`` to merged cells on its logical grid.

    Returns ``(cells, n_rows, n_cols)``. A ``vMerge`` continuation is folded
    into the cell above it that starts at the same grid offset (Word's own
    rule); its row is then covered by that cell's ``row_span``. A continuation
    with nothing to continue is kept as a cell of its own rather than dropped.
    """
    raw = _raw_cells(table)
    n_rows = 1 + max((c.row for c in raw), default=-1)
    right_edge = max((c.grid_end for c in raw), default=0)
    position = _grid_positions(table, right_edge)
    boundaries = _logical_boundaries(raw, right_edge, position)
    n_cols = len(boundaries) - 1

    def nearest(offset: int) -> int:
        """Index of the kept boundary physically closest to a grid offset."""
        return min(
            range(len(boundaries)),
            key=lambda i: abs(position[boundaries[i]] - position[offset]),
        )

    def span_of(c: _RawCell) -> tuple[int, int]:
        col = min(nearest(c.grid_start), n_cols - 1)
        return col, max(col + 1, nearest(c.grid_end))

    placed: list[_Placed] = []
    # Open vertical merges, keyed by grid start offset.
    open_merges: dict[int, _Placed] = {}
    for c in raw:
        if c.vmerge == "continue" and c.grid_start in open_merges:
            open_merges[c.grid_start].row_end = c.row + 1
            continue
        col, col_end = span_of(c)
        entry = _Placed(c.row, c.row + 1, col, col_end, c.text)
        placed.append(entry)
        if c.vmerge == "restart":
            open_merges[c.grid_start] = entry
        else:
            open_merges.pop(c.grid_start, None)

    cells = [
        GridCell(
            row=e.row,
            col=e.col,
            row_span=e.row_end - e.row,
            col_span=e.col_end - e.col,
            text=e.text,
        )
        for e in placed
    ]
    return cells, n_rows, n_cols


def logical_grid(table: _Element) -> list[list[str]]:
    """A rectangular ``rows x logical columns`` list of cell text.

    Every position a merged cell covers - across ``gridSpan`` columns and down
    ``vMerge`` rows - holds that cell's text, so a double-period lab block
    reads the same in both of its slots and a break spanning the week reads as
    a break under every day. Positions no cell covers are ``""``. Use
    :func:`logical_cells` when you need to know which positions were merged.
    """
    cells, n_rows, n_cols = logical_cells(table)
    grid = [["" for _ in range(n_cols)] for _ in range(n_rows)]
    for cell in cells:
        for r in range(cell.row, cell.row + cell.row_span):
            for c in range(cell.col, cell.col + cell.col_span):
                grid[r][c] = cell.text
    return grid


def is_sdt_wrapped(element: _Element) -> bool:
    """True if ``element`` sits inside a content control."""
    return any(ancestor.tag == _W_SDT for ancestor in element.iterancestors())


def physical_grid_width(table: _Element) -> int:
    """Number of ``<w:gridCol>`` entries the table declares."""
    return len(table.findall(f"{qn('w:tblGrid')}/{qn('w:gridCol')}"))
