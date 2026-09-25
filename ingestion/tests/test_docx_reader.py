"""Tests for the shared .docx reader and the table survey.

Two kinds:

* **Tripwires on the real files.** The exact table counts (21/19/15/32) fail
  the moment someone "simplifies" :func:`walk_body` back to
  ``Document.tables``, which sees 4/2/4/32. The grid-shape and cell-text
  assertions fail if merge resolution or line handling regresses.
* **Synthetic XML.** Small hand-written bodies pin down each structural rule -
  content controls, ``gridSpan``, ``vMerge``, sliver folding - in isolation.
"""

from __future__ import annotations

import pytest
from docx.oxml import parse_xml

from ingestion.docx_reader import (
    SLIVER_TWIPS,
    logical_cells,
    logical_grid,
    open_document,
    paragraph_text,
    physical_grid_width,
    walk_body,
)
from ingestion.survey import (
    DEFAULT_REPORT,
    REAL_DIR,
    SOURCE_FILES,
    format_report,
    survey_file,
)

EVEN = "CE_Class_Time_Table-EVEN_sem_25-26.docx"
ODD = "CE_Class_Time_Table-ODD_sem_25-26.docx"
LAB = "Lab CE_Time_Table-Even_sem_25-26.docx"
FACULTY = "Faculty Individual timetable Even 2025-26.docx"


def _tables(name: str) -> list:
    return [el for kind, el in walk_body(open_document(REAL_DIR / name)) if kind == "table"]


# --------------------------------------------------------------------------- #
# Real files: tripwires
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    ("name", "total", "sdt_wrapped"),
    [(EVEN, 21, 17), (ODD, 19, 17), (LAB, 15, 11), (FACULTY, 32, 0)],
)
def test_every_table_is_found(name: str, total: int, sdt_wrapped: int) -> None:
    survey = survey_file(REAL_DIR / name)
    assert len(survey.tables) == total
    assert survey.sdt_wrapped == sdt_wrapped


# The division timetables in the EVEN file; the rest are initials legends.
EVEN_GRIDS = {
    0: "SE-Comp A",
    2: "SE-Comp B",
    4: "SE-Comp C",
    6: "SE-Comp D",
    8: "TE-Comp A",
    10: "TE Comp- B",
    12: "TE Comp - C",
    14: "TE Comp - D",
    17: "BE Comp- A & B",
    19: "M.Tech Comp",
}


def test_even_tables_pair_with_their_division_in_document_order() -> None:
    tables = survey_file(REAL_DIR / EVEN).tables
    for index, division in EVEN_GRIDS.items():
        assert tables[index].header is not None
        assert tables[index].header.endswith(division)
        assert not tables[index].header_shared


def test_even_class_grids_are_six_logical_columns() -> None:
    tables = _tables(EVEN)
    for index in EVEN_GRIDS:
        grid = logical_grid(tables[index])
        assert {len(row) for row in grid} == {6}, index
        assert [cell.strip() for cell in grid[0]] == [
            "Time",
            "Monday",
            "Tuesday",
            "Wednesday",
            "Thursday",
            "Friday",
        ], index


def test_even_physical_grids_really_are_wider() -> None:
    # The premise of the logical grid: 42 grid columns drawn for 6 real ones.
    # (Direct children only - each tblGrid also nests a <w:tblGridChange>
    # revision record holding the previous grid.)
    table = _tables(EVEN)[0]
    _, _, n_cols = logical_cells(table)
    assert physical_grid_width(table) == 42
    assert n_cols == 6


def test_lab_block_cell_keeps_its_four_batch_lines() -> None:
    grid = logical_grid(_tables(EVEN)[0])  # SE-Comp A
    expected = (
        "OS / A / AGN / 702-A\nDAA / B / NR / 603-2\nCCN / C / JS / 606-4\nPCS / D / DN / 608"
    )
    assert grid[4][0] == "11.15-12.15"
    assert grid[4][1] == expected
    # vMerge: the double period reads the same in its second slot.
    assert grid[5][0] == "12.15-1.15"
    assert grid[5][1] == expected


def test_be_sliver_column_folds_without_misplacing_cells() -> None:
    grid = logical_grid(_tables(EVEN)[17])
    # Row 4's Tuesday cell overhangs into a 105-twip sliver of Wednesday; it
    # must stay under Tuesday, and Wednesday's own cell under Wednesday.
    assert grid[4][2].startswith("BDAV /601")
    assert grid[4][3].startswith("BDAV /605")
    assert grid[4][4].startswith("NLP/ RK/ 703")


def test_odd_sub_day_columns_are_not_collapsed() -> None:
    # SE-Comp A (ODD) splits Tuesday and Wednesday into two parallel columns
    # in the afternoon. Both halves are real and must survive.
    grid = logical_grid(_tables(ODD)[0])
    assert len(grid[0]) == 8
    assert grid[0][2] == grid[0][3] == "Tuesday"
    assert grid[0][4] == grid[0][5] == "Wednesday"
    assert grid[6][2] != grid[6][3]


def test_lab_room_number_sits_above_the_utilisation_line() -> None:
    first = survey_file(REAL_DIR / LAB).tables[0]
    assert first.header is not None
    assert first.header.startswith("Lab Utilization")
    assert first.above is not None
    assert first.above.endswith("ROOM NO: 607-B")


def test_faculty_tables_pair_with_a_name() -> None:
    tables = survey_file(REAL_DIR / FACULTY).tables
    assert tables[0].header is not None
    assert tables[0].header.endswith("Prasenjit Bhavathankar")
    assert tables[31].header is not None
    assert tables[31].header.endswith("Dayanand Ambawade")
    assert all(not t.header_shared for t in tables)


def test_committed_survey_report_is_current() -> None:
    report = format_report([survey_file(REAL_DIR / name) for name in SOURCE_FILES])
    assert DEFAULT_REPORT.read_text(encoding="utf-8") == report


# --------------------------------------------------------------------------- #
# Synthetic XML
# --------------------------------------------------------------------------- #

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _p(text: str) -> str:
    return f"<w:p><w:r><w:t xml:space='preserve'>{text}</w:t></w:r></w:p>"


def _tc(text: str, span: int = 1, vmerge: str | None = None) -> str:
    props = ""
    if span > 1:
        props += f'<w:gridSpan w:val="{span}"/>'
    if vmerge == "restart":
        props += '<w:vMerge w:val="restart"/>'
    elif vmerge == "continue":
        props += "<w:vMerge/>"
    paragraphs = "".join(_p(line) for line in text.split("\n"))
    return f"<w:tc><w:tcPr>{props}</w:tcPr>{paragraphs}</w:tc>"


def _tbl(widths: list[int], rows: list[list[str]]) -> str:
    grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
    body = "".join(f"<w:tr>{''.join(r)}</w:tr>" for r in rows)
    return f"<w:tbl><w:tblGrid>{grid}</w:tblGrid>{body}</w:tbl>"


def _sdt(inner: str) -> str:
    return f"<w:sdt><w:sdtPr/><w:sdtContent>{inner}</w:sdtContent></w:sdt>"


class _FakeDocument:
    """Just enough of ``docx.document.Document`` for :func:`walk_body`."""

    def __init__(self, body_xml: str) -> None:
        self.element = parse_xml(f"<w:document {W}><w:body>{body_xml}</w:body></w:document>")


def _table(xml: str):
    return parse_xml(xml.replace("<w:tbl>", f"<w:tbl {W}>", 1))


def test_walk_body_descends_into_content_controls_in_order() -> None:
    one_cell = _tbl([1000], [[_tc("x")]])
    doc = _FakeDocument(
        _p("A") + one_cell + _sdt(_p("B") + one_cell + _sdt(_p("C") + one_cell)) + _p("D")
    )
    seen = [
        (kind, paragraph_text(el) if kind == "paragraph" else "T")
        for kind, el in walk_body(doc)  # type: ignore[arg-type]
    ]
    assert seen == [
        ("paragraph", "A"),
        ("table", "T"),
        ("paragraph", "B"),
        ("table", "T"),
        ("paragraph", "C"),
        ("table", "T"),
        ("paragraph", "D"),
    ]


def test_walk_body_on_empty_body() -> None:
    assert list(walk_body(_FakeDocument(""))) == []  # type: ignore[arg-type]


def test_grid_span_places_cells_at_logical_columns() -> None:
    # Six grid columns, three logical ones: [0,2) [2,5) [5,6).
    table = _table(
        _tbl(
            [1000] * 6,
            [
                [_tc("h0", 2), _tc("h1", 3), _tc("h2")],
                [_tc("a", 2), _tc("b", 3), _tc("c")],
                [_tc("wide", 5), _tc("z")],
            ],
        )
    )
    assert logical_grid(table) == [
        ["h0", "h1", "h2"],
        ["a", "b", "c"],
        ["wide", "wide", "z"],
    ]


def test_identical_text_in_distinct_columns_is_not_collapsed() -> None:
    table = _table(_tbl([1000] * 3, [[_tc("same"), _tc("same"), _tc("same")]]))
    assert logical_grid(table) == [["same", "same", "same"]]


def test_vmerge_fills_down_and_records_row_span() -> None:
    table = _table(
        _tbl(
            [1000, 1000],
            [
                [_tc("t1"), _tc("lab\nblock", vmerge="restart")],
                [_tc("t2"), _tc("", vmerge="continue")],
                [_tc("t3"), _tc("after")],
            ],
        )
    )
    assert logical_grid(table) == [
        ["t1", "lab\nblock"],
        ["t2", "lab\nblock"],
        ["t3", "after"],
    ]
    cells, n_rows, n_cols = logical_cells(table)
    assert (n_rows, n_cols) == (3, 2)
    merged = next(c for c in cells if c.text == "lab\nblock")
    assert (merged.row, merged.col, merged.row_span, merged.col_span) == (0, 1, 2, 1)


def test_orphan_vmerge_continuation_is_kept_not_dropped() -> None:
    table = _table(_tbl([1000], [[_tc("stray", vmerge="continue")]]))
    assert logical_grid(table) == [["stray"]]


def test_sliver_column_folds_but_narrow_real_column_survives() -> None:
    sliver = SLIVER_TWIPS // 3
    # Header edges at 0|1|3; one body row starts a cell at offset 2, which is
    # only a sliver right of offset 1.
    table = _table(
        _tbl(
            [2000, sliver, 2000],
            [
                [_tc("L"), _tc("R", 2)],
                [_tc("left", 2), _tc("right")],
            ],
        )
    )
    assert logical_grid(table) == [["L", "R"], ["left", "right"]]

    real = SLIVER_TWIPS * 2
    table = _table(
        _tbl(
            [2000, real, 2000],
            [
                [_tc("L"), _tc("R", 2)],
                [_tc("left", 2), _tc("right")],
            ],
        )
    )
    assert logical_grid(table) == [["L", "R", "R"], ["left", "left", "right"]]
